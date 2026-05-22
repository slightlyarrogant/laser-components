#!/usr/bin/env python3
"""
Generate AI Sales Intelligence Platform PowerPoint Presentation
Converts the HTML presentation to a professional .pptx file.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# ─── Constants ───────────────────────────────────────────────────────────────

BLUE = RGBColor(0x4A, 0xBA, 0xFF)
PURPLE = RGBColor(0x9B, 0x6D, 0xFF)
DARK = RGBColor(0x1A, 0x1A, 0x2E)
DARK_BG = RGBColor(0x16, 0x16, 0x2B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF8, 0xF9, 0xFA)
MID_GRAY = RGBColor(0x88, 0x88, 0x88)
TEXT_GRAY = RGBColor(0x66, 0x66, 0x66)
TEXT_DARK = RGBColor(0x44, 0x44, 0x44)
GREEN = RGBColor(0x22, 0xC5, 0x5E)
RED = RGBColor(0xEF, 0x44, 0x44)
TABLE_ALT = RGBColor(0xF0, 0xF4, 0xFF)
HIGHLIGHT_BG = RGBColor(0xE8, 0xF0, 0xFF)

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

FONT_TITLE = "Calibri"
FONT_BODY = "Calibri"

LOGO_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "logo.png")
OUTPUT_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                           "ai-sales-platform-presentation.pptx")


# ─── Helpers ─────────────────────────────────────────────────────────────────

def set_slide_bg(slide, color):
    """Set solid background color for a slide."""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_accent_bar(slide, left, top, width, height, color):
    """Add a colored rectangle as an accent bar."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_textbox(slide, left, top, width, height, text, font_size=18,
                color=DARK, bold=False, alignment=PP_ALIGN.LEFT,
                font_name=FONT_BODY):
    """Add a text box with specified formatting."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_bullet_list(slide, left, top, width, height, items, font_size=16,
                    color=DARK, bullet_color=BLUE, spacing=Pt(8)):
    """Add a bulleted list to the slide."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True

    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = FONT_BODY
        p.space_after = spacing
        p.level = 0
        # Use bullet character
        p.text = "\u2022  " + item
    return txBox


def add_gradient_header_bar(slide):
    """Add blue accent bar at top of content slides."""
    # Top gradient bar (blue)
    add_accent_bar(slide, Inches(0), Inches(0), SLIDE_WIDTH, Inches(0.06), BLUE)
    # Secondary thin purple bar
    add_accent_bar(slide, Inches(0), Inches(0.06), Inches(6), Inches(0.03), PURPLE)


def add_slide_number(slide, number, total=15):
    """Add slide number in bottom right."""
    add_textbox(slide, Inches(11.5), Inches(7.0), Inches(1.5), Inches(0.4),
                f"{number} / {total}", font_size=11, color=MID_GRAY,
                alignment=PP_ALIGN.RIGHT)


def add_section_label(slide, text, left=Inches(0.8), top=Inches(0.5)):
    """Add a blue/purple section label above the title."""
    tb = add_textbox(slide, left, top, Inches(6), Inches(0.5),
                     text, font_size=16, color=BLUE, bold=True)
    return tb


def add_slide_title(slide, text, left=Inches(0.8), top=Inches(1.0)):
    """Add the main slide title."""
    tb = add_textbox(slide, left, top, Inches(11), Inches(0.8),
                     text, font_size=30, color=DARK, bold=True)
    return tb


def add_highlight_box(slide, text, left, top, width, height,
                      bg_color=BLUE, text_color=WHITE, font_size=16):
    """Add a colored highlight box with text."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.fill.background()

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.15)
    tf.margin_bottom = Inches(0.15)
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = text_color
    p.font.name = FONT_BODY
    p.alignment = PP_ALIGN.CENTER
    return shape


def add_card(slide, left, top, width, height, title, body_lines,
             title_color=PURPLE, bg_color=LIGHT_GRAY):
    """Add a card-style box with title and body text."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.fill.background()

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.2)
    tf.margin_bottom = Inches(0.15)

    # Title
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(16)
    p.font.color.rgb = title_color
    p.font.bold = True
    p.font.name = FONT_BODY
    p.space_after = Pt(8)

    # Body lines
    for line in body_lines:
        p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_GRAY
        p.font.name = FONT_BODY
        p.space_after = Pt(4)

    return shape


def add_table(slide, left, top, width, height, headers, rows,
              header_color=BLUE, header_text_color=WHITE):
    """Add a styled table."""
    n_rows = len(rows) + 1
    n_cols = len(headers)
    table_shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = table_shape.table

    # Set column widths evenly
    col_width = int(width / n_cols)
    for i in range(n_cols):
        table.columns[i].width = col_width

    # Header row
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_color
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.size = Pt(14)
            paragraph.font.color.rgb = header_text_color
            paragraph.font.bold = True
            paragraph.font.name = FONT_BODY

    # Data rows
    for r, row_data in enumerate(rows):
        for c, cell_text in enumerate(row_data):
            cell = table.cell(r + 1, c)
            cell.text = str(cell_text)
            # Alternate row background
            if r % 2 == 1:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT_GRAY
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(13)
                paragraph.font.color.rgb = DARK
                paragraph.font.name = FONT_BODY

    return table_shape


def add_phase_circle(slide, left, top, number, title, desc_lines):
    """Add a numbered phase circle with title and description."""
    # Circle
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, Inches(0.7), Inches(0.7))
    circle.fill.solid()
    circle.fill.fore_color.rgb = BLUE
    circle.line.fill.background()
    tf = circle.text_frame
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    p = tf.paragraphs[0]
    p.text = str(number)
    p.font.size = Pt(22)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.font.name = FONT_BODY
    p.alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    # Title
    add_textbox(slide, left - Inches(0.3), top + Inches(0.85), Inches(1.3), Inches(0.4),
                title, font_size=14, color=DARK, bold=True, alignment=PP_ALIGN.CENTER)

    # Description
    desc_text = "\n".join(desc_lines)
    add_textbox(slide, left - Inches(0.3), top + Inches(1.2), Inches(1.3), Inches(0.6),
                desc_text, font_size=11, color=TEXT_GRAY, alignment=PP_ALIGN.CENTER)


def add_arrow(slide, left, top):
    """Add a small arrow between phases."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, Inches(0.5), Inches(0.3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = BLUE
    shape.line.fill.background()


# ─── Slide Builders ──────────────────────────────────────────────────────────

def build_slide_01_title(prs):
    """Slide 1: Title slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    set_slide_bg(slide, DARK)

    # Top accent bar
    add_accent_bar(slide, Inches(0), Inches(0), SLIDE_WIDTH, Inches(0.08), BLUE)

    # Logo
    if os.path.exists(LOGO_PATH):
        slide.shapes.add_picture(LOGO_PATH, Inches(5.4), Inches(1.0),
                                 height=Inches(1.5))

    # Main title
    add_textbox(slide, Inches(1.5), Inches(2.8), Inches(10.3), Inches(1.2),
                "AI Sales Intelligence Platform",
                font_size=42, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

    # Subtitle
    add_textbox(slide, Inches(2.5), Inches(4.1), Inches(8.3), Inches(0.6),
                "For Laser Components B2B Sales Research",
                font_size=22, color=RGBColor(0xBB, 0xBB, 0xDD), bold=False,
                alignment=PP_ALIGN.CENTER)

    # Company info
    add_textbox(slide, Inches(3.5), Inches(5.0), Inches(6.3), Inches(0.5),
                "AUTOOFFICE Sp. z o.o.  |  January 2026",
                font_size=16, color=MID_GRAY, alignment=PP_ALIGN.CENTER)

    # Bottom accent bar
    add_accent_bar(slide, Inches(4), Inches(6.3), Inches(5.3), Inches(0.04), BLUE)

    add_slide_number(slide, 1)


def build_slide_02_challenge(prs):
    """Slide 2: The Challenge."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_gradient_header_bar(slide)

    add_section_label(slide, "The Challenge")
    add_slide_title(slide, "Current Pain Points in Sales Research")

    items = [
        "Generic AI tools don't understand your products",
        "Research results lost in chat history",
        "Manual verification of supply chain position",
        "No knowledge accumulation between sessions",
        "Each researcher starts from scratch",
        "Inconsistent methodology across team",
    ]
    add_bullet_list(slide, Inches(1.2), Inches(2.2), Inches(10), Inches(4.5),
                    items, font_size=18, color=DARK)

    # Decorative side bar
    add_accent_bar(slide, Inches(0.8), Inches(2.2), Inches(0.06), Inches(4.5), BLUE)

    add_slide_number(slide, 2)


def build_slide_03_solution(prs):
    """Slide 3: The Solution."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_gradient_header_bar(slide)

    add_section_label(slide, "The Solution")
    add_slide_title(slide, "AI-Powered Sales Research Platform")

    items = [
        "Custom MCP Server with 28+ specialized tools",
        "Claude AI Integration (Anthropic's state-of-the-art models)",
        "Domain-Specific Skills with photonics industry knowledge",
        "Persistent database for leads and research history",
        "Scalable architecture from 1 user to enterprise",
        "Learning system that improves with use",
    ]
    add_bullet_list(slide, Inches(1.2), Inches(2.2), Inches(10), Inches(4.5),
                    items, font_size=18, color=DARK)

    add_accent_bar(slide, Inches(0.8), Inches(2.2), Inches(0.06), Inches(4.5), PURPLE)

    add_slide_number(slide, 3)


def build_slide_04_differentiators(prs):
    """Slide 4: Key Differentiators."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_gradient_header_bar(slide)

    add_section_label(slide, "Key Differentiators")
    add_slide_title(slide, "Why This Platform?")

    headers = ["Generic AI (ChatGPT, etc.)", "This Platform"]
    rows = [
        ["No product knowledge", "245+ products in database"],
        ["Text output to copy/paste", "Leads stored, tagged, exportable"],
        ["Each session starts fresh", "Learning scratch pad"],
        ["Generic B2B understanding", "Supply chain positioning built-in"],
        ["No actions, just answers", "Creates leads, exports, reports"],
    ]
    tbl = add_table(slide, Inches(1.0), Inches(2.2), Inches(11.3), Inches(4.0),
                    headers, rows)

    # Bold the second column values
    table = tbl.table
    for r in range(1, len(rows) + 1):
        cell = table.cell(r, 1)
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.bold = True
            paragraph.font.color.rgb = PURPLE

    add_slide_number(slide, 4)


def build_slide_05_supply_chain(prs):
    """Slide 5: Supply Chain Intelligence."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_gradient_header_bar(slide)

    add_section_label(slide, "Supply Chain Intelligence")
    add_slide_title(slide, "Targeting the Right Companies")

    # Supply chain flow box
    flow_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                        Inches(0.8), Inches(2.2),
                                        Inches(11.7), Inches(1.3))
    flow_shape.fill.solid()
    flow_shape.fill.fore_color.rgb = LIGHT_GRAY
    flow_shape.line.color.rgb = BLUE
    flow_shape.line.width = Pt(2)

    tf = flow_shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.15)

    p = tf.paragraphs[0]
    p.text = "[Laser Components]  -->  [Component Integrators]  -->  [Module Makers]  -->  [System Integrators]  -->  [OEMs]"
    p.font.size = Pt(15)
    p.font.name = "Consolas"
    p.font.color.rgb = DARK

    p2 = tf.add_paragraph()
    p2.text = "                                              TARGETS                                             TARGETS"
    p2.font.size = Pt(13)
    p2.font.name = "Consolas"
    p2.font.color.rgb = BLUE
    p2.font.bold = True

    # Correct targets card
    add_card(slide, Inches(0.8), Inches(3.8), Inches(5.5), Inches(3.0),
             "CORRECT TARGETS",
             [
                 "\u2022  Board manufacturers",
                 "\u2022  Optical assembly houses",
                 "\u2022  Contract manufacturers",
                 "\u2022  Module builders",
             ],
             title_color=GREEN)

    # Wrong targets card
    add_card(slide, Inches(7.0), Inches(3.8), Inches(5.5), Inches(3.0),
             "WRONG TARGETS",
             [
                 "\u2022  Final OEMs (BMW, Tesla)",
                 "\u2022  System integrators (Velodyne)",
                 "\u2022  Consumer brands",
                 "\u2022  Retailers",
             ],
             title_color=RED)

    add_slide_number(slide, 5)


def build_slide_06_domain_knowledge(prs):
    """Slide 6: Embedded Domain Knowledge."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_gradient_header_bar(slide)

    add_section_label(slide, "Embedded Domain Knowledge")
    add_slide_title(slide, "Sales Research Skill Includes:")

    # Industry Intelligence card
    add_card(slide, Inches(0.8), Inches(2.2), Inches(5.5), Inches(3.0),
             "Industry Intelligence",
             [
                 "\u2022  Product-application mapping",
                 "\u2022  Wavelength correlations (905nm -> LIDAR)",
                 "\u2022  European photonics clusters",
                 "\u2022  Trade show calendar",
             ],
             title_color=PURPLE)

    # Automated Workflows card
    add_card(slide, Inches(7.0), Inches(2.2), Inches(5.5), Inches(3.0),
             "Automated Workflows",
             [
                 "\u2022  Geographic prospecting",
                 "\u2022  Application discovery",
                 "\u2022  Lead enrichment & scoring",
                 "\u2022  Competitive analysis",
             ],
             title_color=PURPLE)

    # Highlight box
    add_highlight_box(slide,
                      "Built-in lead qualification: supply chain position, technical fit, company viability",
                      Inches(0.8), Inches(5.7), Inches(11.7), Inches(0.8),
                      bg_color=BLUE, text_color=WHITE, font_size=16)

    add_slide_number(slide, 6)


def build_slide_07_tools(prs):
    """Slide 7: 28 MCP Tools."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_gradient_header_bar(slide)

    add_section_label(slide, "Comprehensive Toolset")
    add_slide_title(slide, "28 MCP Tools")

    # Three tool category cards
    categories = [
        ("15", "Core Data Tools", "Products, Leads, Applications,\nCategories CRUD"),
        ("6", "AI Integration", "Market analysis, enrichment,\nscoring, insights"),
        ("7", "Analytics & Workflow", "Batch ops, exports, reports,\nactivity feed"),
    ]

    for i, (num, label, desc) in enumerate(categories):
        left = Inches(0.8 + i * 4.1)
        # Card background
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                       left, Inches(2.2), Inches(3.6), Inches(3.0))
        shape.fill.solid()
        shape.fill.fore_color.rgb = LIGHT_GRAY
        shape.line.fill.background()

        # Big number
        add_textbox(slide, left, Inches(2.5), Inches(3.6), Inches(1.0),
                    num, font_size=48, color=BLUE, bold=True,
                    alignment=PP_ALIGN.CENTER)

        # Label
        add_textbox(slide, left, Inches(3.4), Inches(3.6), Inches(0.5),
                    label, font_size=16, color=DARK, bold=True,
                    alignment=PP_ALIGN.CENTER)

        # Description
        add_textbox(slide, left, Inches(3.9), Inches(3.6), Inches(0.8),
                    desc, font_size=12, color=TEXT_GRAY,
                    alignment=PP_ALIGN.CENTER)

    # Highlight box at bottom
    add_highlight_box(slide,
                      "All tools accessible via natural language queries - no coding required",
                      Inches(0.8), Inches(5.7), Inches(11.7), Inches(0.8),
                      bg_color=BLUE, text_color=WHITE, font_size=16)

    add_slide_number(slide, 7)


def build_slide_08_tech_stack(prs):
    """Slide 8: Technology Stack."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_gradient_header_bar(slide)

    add_section_label(slide, "Enterprise-Grade Foundation")
    add_slide_title(slide, "Technology Stack")

    tech_items = [
        ("AI Core", "Anthropic Claude Opus 4.5 (SOTA reasoning)"),
        ("Protocol", "MCP - Anthropic's official standard"),
        ("Database", "PostgreSQL + Prisma ORM"),
        ("Research", "Perplexity AI integration"),
        ("Security", "Data on YOUR infrastructure"),
        ("Compliance", "GDPR compliant by design"),
    ]

    for i, (label, desc) in enumerate(tech_items):
        row = i // 2
        col = i % 2
        left = Inches(0.8 + col * 6.2)
        top = Inches(2.2 + row * 1.15)

        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                       left, top, Inches(5.7), Inches(0.9))
        shape.fill.solid()
        shape.fill.fore_color.rgb = LIGHT_GRAY
        shape.line.fill.background()

        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.15)
        p = tf.paragraphs[0]

        # Bold label
        run_label = p.add_run()
        run_label.text = label + "   "
        run_label.font.size = Pt(15)
        run_label.font.color.rgb = PURPLE
        run_label.font.bold = True
        run_label.font.name = FONT_BODY

        # Description
        run_desc = p.add_run()
        run_desc.text = desc
        run_desc.font.size = Pt(14)
        run_desc.font.color.rgb = DARK
        run_desc.font.name = FONT_BODY

    # Highlight box
    add_highlight_box(slide,
                      "Anthropic: $60B+ valuation, leading AI safety company, trusted by Amazon, Google, Salesforce",
                      Inches(0.8), Inches(5.9), Inches(11.7), Inches(0.8),
                      bg_color=BLUE, text_color=WHITE, font_size=14)

    add_slide_number(slide, 8)


def build_slide_09_interface(prs):
    """Slide 9: Interface Evolution."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_gradient_header_bar(slide)

    add_section_label(slide, "Interface Evolution")
    add_slide_title(slide, "From CLI to Web")

    # Current: CLI card
    shape1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                    Inches(0.8), Inches(2.2), Inches(5.7), Inches(4.5))
    shape1.fill.solid()
    shape1.fill.fore_color.rgb = LIGHT_GRAY
    shape1.line.fill.background()

    tf1 = shape1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.3)
    tf1.margin_top = Inches(0.25)

    p = tf1.paragraphs[0]
    p.text = "Current: Claude Code (CLI)"
    p.font.size = Pt(18)
    p.font.color.rgb = PURPLE
    p.font.bold = True
    p.font.name = FONT_BODY
    p.space_after = Pt(16)

    for item in ["Full power for tech users", "Scriptable automation", "Fast execution"]:
        p = tf1.add_paragraph()
        p.text = "\u2022  " + item
        p.font.size = Pt(14)
        p.font.color.rgb = DARK
        p.font.name = FONT_BODY
        p.space_after = Pt(6)

    # CLI demo box
    p = tf1.add_paragraph()
    p.space_before = Pt(12)
    p.space_after = Pt(0)

    cli_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                     Inches(1.1), Inches(4.6), Inches(5.1), Inches(1.2))
    cli_box.fill.solid()
    cli_box.fill.fore_color.rgb = DARK
    cli_box.line.fill.background()

    tf_cli = cli_box.text_frame
    tf_cli.word_wrap = True
    tf_cli.margin_left = Inches(0.2)
    tf_cli.margin_top = Inches(0.15)

    p = tf_cli.paragraphs[0]
    p.text = "$ claude"
    p.font.size = Pt(12)
    p.font.color.rgb = BLUE
    p.font.name = "Consolas"

    p2 = tf_cli.add_paragraph()
    p2.text = "> Find Czech rangefinder companies"
    p2.font.size = Pt(12)
    p2.font.color.rgb = WHITE
    p2.font.name = "Consolas"

    p3 = tf_cli.add_paragraph()
    p3.text = "[Creates 9 leads in database]"
    p3.font.size = Pt(12)
    p3.font.color.rgb = GREEN
    p3.font.name = "Consolas"

    # Coming Q2 2026 card
    shape2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                    Inches(7.0), Inches(2.2), Inches(5.7), Inches(4.5))
    shape2.fill.solid()
    shape2.fill.fore_color.rgb = LIGHT_GRAY
    shape2.line.fill.background()

    tf2 = shape2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = Inches(0.3)
    tf2.margin_top = Inches(0.25)

    p = tf2.paragraphs[0]
    p.text = "Coming Q2 2026: Claude Cowork"
    p.font.size = Pt(18)
    p.font.color.rgb = PURPLE
    p.font.bold = True
    p.font.name = FONT_BODY
    p.space_after = Pt(16)

    for item in ["Web-based interface", "Team collaboration",
                 "Visual lead management", "Non-technical user friendly"]:
        p = tf2.add_paragraph()
        p.text = "\u2022  " + item
        p.font.size = Pt(14)
        p.font.color.rgb = DARK
        p.font.name = FONT_BODY
        p.space_after = Pt(6)

    p = tf2.add_paragraph()
    p.space_before = Pt(20)
    run = p.add_run()
    run.text = "Same backend"
    run.font.size = Pt(14)
    run.font.color.rgb = DARK
    run.font.bold = True
    run.font.name = FONT_BODY

    run2 = p.add_run()
    run2.text = " - zero migration cost"
    run2.font.size = Pt(14)
    run2.font.color.rgb = TEXT_GRAY
    run2.font.name = FONT_BODY

    add_slide_number(slide, 9)


def build_slide_10_deployment(prs):
    """Slide 10: Deployment Options."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_gradient_header_bar(slide)

    add_section_label(slide, "Flexible Architecture")
    add_slide_title(slide, "Deployment Options")

    options = [
        ("Local", ["Single laptop", "Zero infrastructure", "Full control"]),
        ("Server", ["On-premise Linux", "Team access via VPN", "Shared database"]),
        ("Cloud", ["AWS / Azure / GCP", "Access from anywhere", "99.9% uptime"]),
    ]

    for i, (title, lines) in enumerate(options):
        left = Inches(0.8 + i * 4.1)

        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                       left, Inches(2.2), Inches(3.6), Inches(3.2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = LIGHT_GRAY
        shape.line.fill.background()

        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.25)
        tf.vertical_anchor = MSO_ANCHOR.TOP

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(22)
        p.font.color.rgb = BLUE
        p.font.bold = True
        p.font.name = FONT_BODY
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(16)

        for line in lines:
            p = tf.add_paragraph()
            p.text = line
            p.font.size = Pt(14)
            p.font.color.rgb = TEXT_GRAY
            p.font.name = FONT_BODY
            p.alignment = PP_ALIGN.CENTER
            p.space_after = Pt(8)

    # Highlight box
    add_highlight_box(slide,
                      "Database flexibility: Local  -->  Company Server  -->  Managed Cloud (Supabase, RDS)",
                      Inches(0.8), Inches(5.9), Inches(11.7), Inches(0.8),
                      bg_color=BLUE, text_color=WHITE, font_size=14)

    add_slide_number(slide, 10)


def build_slide_11_pricing(prs):
    """Slide 11: Pricing Overview."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_gradient_header_bar(slide)

    add_section_label(slide, "Investment")
    add_slide_title(slide, "Pricing Tiers")

    headers = ["Tier", "Users", "Setup", "Monthly", "Year 1 Total"]
    rows = [
        ["Personal", "1", "\u20ac4,500", "\u20ac500", "\u20ac10,500"],
        ["Team", "Up to 10", "\u20ac6,000", "\u20ac3,500", "\u20ac48,000"],
        ["Enterprise", "Up to 50", "\u20ac12,000", "\u20ac12,000", "\u20ac156,000"],
        ["Strategic", "Unlimited", "Custom", "\u20ac16,000+", "\u20ac200,000+"],
    ]
    tbl = add_table(slide, Inches(0.8), Inches(2.2), Inches(11.7), Inches(3.5),
                    headers, rows)

    # Bold tier names
    table = tbl.table
    for r in range(1, len(rows) + 1):
        cell = table.cell(r, 0)
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.bold = True

    # Footer note
    add_textbox(slide, Inches(0.8), Inches(6.1), Inches(11), Inches(0.5),
                "All tiers include: API costs (Claude + Perplexity), maintenance, support, Claude Cowork migration",
                font_size=13, color=TEXT_GRAY)

    add_slide_number(slide, 11)


def build_slide_12_roi(prs):
    """Slide 12: ROI Comparison."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_gradient_header_bar(slide)

    add_section_label(slide, "Return on Investment")
    add_slide_title(slide, "vs. Alternatives (10 users)")

    headers = ["Solution", "Annual Cost", "What's Missing"]
    rows = [
        ["ZoomInfo", "\u20ac150,000+", "Database only, no AI, no product context"],
        ["LinkedIn Sales Nav", "\u20ac15,000", "Contacts only, no lead storage"],
        ["Dedicated FTE", "\u20ac60,000+", "1 person, knowledge leaves when they do"],
        ["Custom AI Build", "\u20ac200,000+", "One-time, no maintenance included"],
        ["This Platform (Team)", "\u20ac48,000", "Full capability + APIs + support included"],
    ]
    tbl = add_table(slide, Inches(0.8), Inches(2.2), Inches(11.7), Inches(3.8),
                    headers, rows)

    # Highlight the last row (our platform)
    table = tbl.table
    for c in range(3):
        cell = table.cell(5, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = HIGHLIGHT_BG
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.bold = True
            paragraph.font.color.rgb = BLUE

    # Highlight box
    add_highlight_box(slide,
                      "3x cheaper than ZoomInfo with domain-specific AI + all API costs included",
                      Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.7),
                      bg_color=BLUE, text_color=WHITE, font_size=15)

    add_slide_number(slide, 12)


def build_slide_13_implementation(prs):
    """Slide 13: Implementation Roadmap."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_gradient_header_bar(slide)

    add_section_label(slide, "Implementation Roadmap")
    add_slide_title(slide, "3-Phase Approach")

    phases = [
        (1, "Pilot", ["Month 1-3", "Setup, training", "2-3 test markets"]),
        (2, "Expansion", ["Month 4-6", "Full team rollout", "Integrations"]),
        (3, "Cowork", ["Q2-Q3 2026", "Web interface", "Non-technical users"]),
    ]

    for i, (num, title, desc) in enumerate(phases):
        left = Inches(1.5 + i * 4.0)
        add_phase_circle(slide, left, Inches(2.8), num, title, desc)
        if i < len(phases) - 1:
            add_arrow(slide, Inches(3.0 + i * 4.0), Inches(3.0))

    # Highlight box
    add_highlight_box(slide,
                      "Weekly check-ins during pilot  |  Continuous optimization  |  Same backend throughout",
                      Inches(0.8), Inches(5.7), Inches(11.7), Inches(0.8),
                      bg_color=BLUE, text_color=WHITE, font_size=15)

    add_slide_number(slide, 13)


def build_slide_14_next_steps(prs):
    """Slide 14: Next Steps."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_gradient_header_bar(slide)

    add_section_label(slide, "Getting Started")
    add_slide_title(slide, "Next Steps")

    steps = [
        (1, "Discovery\nCall", ["Live demo", "1 hour"]),
        (2, "Pilot\nScope", ["Define test", "markets"]),
        (3, "Contract", ["Service", "agreement"]),
        (4, "Deployment", ["Setup +", "training"]),
        (5, "Go-Live", ["Pilot", "starts"]),
    ]

    for i, (num, title, desc) in enumerate(steps):
        left = Inches(0.7 + i * 2.5)
        add_phase_circle(slide, left, Inches(2.8), num, title, desc)
        if i < len(steps) - 1:
            add_arrow(slide, Inches(1.8 + i * 2.5), Inches(3.0))

    add_slide_number(slide, 14)


def build_slide_15_contact(prs):
    """Slide 15: Contact / Thank You."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK)

    # Top accent bar
    add_accent_bar(slide, Inches(0), Inches(0), SLIDE_WIDTH, Inches(0.08), BLUE)

    # Section label
    add_textbox(slide, Inches(1.5), Inches(0.6), Inches(10.3), Inches(0.6),
                "Let's Talk",
                font_size=20, color=BLUE, bold=True, alignment=PP_ALIGN.CENTER)

    # Logo
    if os.path.exists(LOGO_PATH):
        slide.shapes.add_picture(LOGO_PATH, Inches(5.7), Inches(1.3),
                                 height=Inches(1.2))

    # Name
    add_textbox(slide, Inches(2.5), Inches(2.8), Inches(8.3), Inches(0.6),
                "Bogdan Czarnecki",
                font_size=28, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

    # Email
    add_textbox(slide, Inches(2.5), Inches(3.5), Inches(8.3), Inches(0.4),
                "bogdan.czarnecki@ikrystyna.pl",
                font_size=16, color=BLUE, alignment=PP_ALIGN.CENTER)

    # Phone
    add_textbox(slide, Inches(2.5), Inches(3.95), Inches(8.3), Inches(0.4),
                "+48 668 814 400",
                font_size=16, color=BLUE, alignment=PP_ALIGN.CENTER)

    # Divider
    add_accent_bar(slide, Inches(4.5), Inches(4.6), Inches(4.3), Inches(0.02), MID_GRAY)

    # Company
    add_textbox(slide, Inches(2.5), Inches(4.9), Inches(8.3), Inches(0.4),
                "AUTOOFFICE Sp. z o.o.",
                font_size=16, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

    # Address
    add_textbox(slide, Inches(2.5), Inches(5.35), Inches(8.3), Inches(0.6),
                "Sw. Wojciecha 2/9\n45-015 Opole, Poland",
                font_size=13, color=MID_GRAY, alignment=PP_ALIGN.CENTER)

    # Legal
    add_textbox(slide, Inches(2.0), Inches(6.2), Inches(9.3), Inches(0.4),
                "KRS: 0001117764  |  NIP: 7543370963  |  REGON: 529219565",
                font_size=11, color=RGBColor(0x66, 0x66, 0x77),
                alignment=PP_ALIGN.CENTER)

    # Bottom bar
    add_accent_bar(slide, Inches(0), Inches(7.35), SLIDE_WIDTH, Inches(0.08), BLUE)

    add_slide_number(slide, 15)


# ─── Main ────────────────────────────────────────────────────────────────────

def main():
    prs = Presentation()

    # Set widescreen 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Build all 15 slides
    build_slide_01_title(prs)
    build_slide_02_challenge(prs)
    build_slide_03_solution(prs)
    build_slide_04_differentiators(prs)
    build_slide_05_supply_chain(prs)
    build_slide_06_domain_knowledge(prs)
    build_slide_07_tools(prs)
    build_slide_08_tech_stack(prs)
    build_slide_09_interface(prs)
    build_slide_10_deployment(prs)
    build_slide_11_pricing(prs)
    build_slide_12_roi(prs)
    build_slide_13_implementation(prs)
    build_slide_14_next_steps(prs)
    build_slide_15_contact(prs)

    prs.save(OUTPUT_PATH)
    print(f"Presentation saved to: {OUTPUT_PATH}")
    print(f"Total slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
